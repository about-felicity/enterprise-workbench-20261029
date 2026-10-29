"""Derived illustrated exports; original article receipts remain immutable."""
import hashlib,io,re
from pathlib import Path
from . import runtime,studio,product_images,knowledge_library
from geo_source_cards.runtime import process_lock,atomic_json,atomic_write

VERSION='product-image-v2'
def normalized(text):return re.sub(r'[^\w\u4e00-\u9fff]','',text).casefold().replace('fbcy','')

def select(root,folder):
    job=runtime.read(root/'job.json');pid=job.get('product_id')
    archived=runtime.read(folder/'rewrite_context.json').get('product') or runtime.read(root/'days'/folder.relative_to(root/'approved').parts[0]/'product.json')
    if not archived:archived=runtime.read(root/'product.json')
    products=studio.products()
    # Match the full product identity, never just its brand or category.
    candidates=[p for p in products if normalized(p['product']['name'])==normalized(archived.get('name',''))]
    if len(candidates)!=1:return None
    pid=candidates[0]['product_id']
    libraries={r['library_id'] for r in knowledge_library.records() if r.get('product_id')==pid}
    images=[i for i in product_images.images() if i.get('product_id')==pid or (not i.get('product_id') and i.get('library_id') in libraries)]
    if not images:return None
    return archived,images[0]

def anchor(lines,name):
    name=normalized(name)
    matches=[i for i,s in enumerate(lines) if name and name in normalized(s) and not s.lstrip().startswith('# ')]
    headings=[i for i in matches if re.match(r'^\s*(?:#{2,6}\s|\*\*|\d+[.、)）])',lines[i]) and len(lines[i])<180]
    return (headings or matches or [None])[0]

def project(root,folder):
    root,folder=Path(root),Path(folder)
    chosen=select(root,folder)
    original={'word':folder/'article.docx','markdown':folder/'draft.md','image':None}
    if not chosen:return original
    product,meta=chosen;markdown=(folder/'draft.md').read_text('utf8');lines=markdown.splitlines()
    index=anchor(lines,product['name'])
    if index is None:return original
    raw,mime=product_images.read(meta['image_id'])
    if hashlib.sha256(raw).hexdigest()!=meta['sha256']:raise ValueError('产品图片校验失败，请重新上传')
    from PIL import Image,ImageOps
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    key=runtime.digest([VERSION,str(folder),hashlib.sha256((folder/'article.docx').read_bytes()).hexdigest(),hashlib.sha256(markdown.encode()).hexdigest(),meta['sha256'],product['name']])
    out=runtime.BASE/'illustrated'/key
    with process_lock(out.with_suffix('.lock'),blocking=True):
        if (out/'complete.json').exists():return {'word':out/'article.docx','markdown':out/'draft.md','image':out/'product.png','image_id':meta['image_id']}
        document=Document(folder/'article.docx')
        target=normalized(re.sub(r'^\s*#{1,6}\s*','',lines[index]).replace('**',''))
        matches=[p for p in document.paragraphs if normalized(p.text)==target]
        if len(matches)!=1:return original
        out.mkdir(parents=True,exist_ok=True)
        with Image.open(io.BytesIO(raw)) as source:
            image=ImageOps.exif_transpose(source).convert('RGB' if source.mode not in ('RGBA','LA') else 'RGBA')
            w,h=image.size;image.save(out/'product.png')
        section=document.sections[0]
        max_width=min(14,(section.page_width-section.left_margin-section.right_margin)/360000)
        scale=min(max_width/w,9/h,2.54/96)
        p=document.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(0)
        p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.keep_together=True
        p.add_run().add_picture(str(out/'product.png'),width=Cm(w*scale),height=Cm(h*scale))
        matches[0]._p.addnext(p._p);matches[0].paragraph_format.keep_with_next=True
        staged=out/'pending.docx';document.save(staged);staged.replace(out/'article.docx')
        lines[index:index+1]=[lines[index],'',f'![产品图片](/api/geo-creation/product_image?id={meta["image_id"]})','']
        atomic_write(out/'draft.md','\n'.join(lines)+'\n')
        atomic_json(out/'complete.json',{'version':VERSION,'image_id':meta['image_id'],'source_image_sha256':meta['sha256'],'width_cm':w*scale,'height_cm':h*scale,'anchor':matches[0].text,'original_word_sha256':hashlib.sha256((folder/'article.docx').read_bytes()).hexdigest(),'word_sha256':hashlib.sha256((out/'article.docx').read_bytes()).hexdigest()})
    return {'word':out/'article.docx','markdown':out/'draft.md','image':out/'product.png','image_id':meta['image_id']}
