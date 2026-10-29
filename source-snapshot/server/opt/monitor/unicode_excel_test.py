from openpyxl import Workbook, load_workbook

p = r"C:\Users\AMD\Desktop\monitor\unicode_test_file.xlsx"
s1 = chr(0x89c6) + chr(0x9891)
s2 = chr(0x6309) + chr(0x7c7b) + chr(0x578b) + chr(0x6c47) + chr(0x603b)
wb = Workbook()
ws = wb.active
ws["A1"] = s1
ws["A2"] = s2
wb.save(p)
wb2 = load_workbook(p, read_only=True, data_only=True)
for cell in ["A1", "A2"]:
    v = wb2.active[cell].value
    print(repr(v), [hex(ord(ch)) for ch in v])
