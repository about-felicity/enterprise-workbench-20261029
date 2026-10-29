import csv
import json
import os
import sys
from datetime import datetime
from urllib.parse import urlparse


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_CSV = os.path.join(BASE_DIR, "deepseek_refs_result.csv")
OUT_XLSX = os.path.join(BASE_DIR, "deepseek_refs_result.xlsx")

FIELDS = [
    "run_no",
    "run_time",
    "chat_id",
    "chat_title",
    "question",
    "page_url",
    "status",
    "complete",
    "count",
    "expected_count",
    "index",
    "title",
    "href",
    "source",
    "domain",
    "extracted_at",
]


def load_payload(raw):
    raw = (raw or "").strip()
    if not raw:
        raise ValueError("empty json")
    return json.loads(raw)


def chat_id_from_url(url):
    return (url or "").rstrip("/").split("/")[-1]


def domain_from_url(url):
    return urlparse(url or "").netloc.lower().split(":")[0]


def next_run_no():
    if not os.path.exists(OUT_CSV):
        return 1
    try:
        with open(OUT_CSV, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            nums = []
            for row in reader:
                try:
                    nums.append(int(row.get("run_no") or 0))
                except Exception:
                    pass
        return (max(nums) if nums else 0) + 1
    except Exception:
        return 1


def ensure_csv_header():
    if os.path.exists(OUT_CSV) and os.path.getsize(OUT_CSV) > 0:
        return False
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        csv.DictWriter(f, fieldnames=FIELDS).writeheader()
    return True


def append_csv(payload):
    ensure_csv_header()
    run_no = next_run_no()
    run_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = payload.get("items") or []

    with open(OUT_CSV, "a", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        for item in rows:
            href = item.get("href") or ""
            writer.writerow({
                "run_no": run_no,
                "run_time": run_time,
                "chat_id": chat_id_from_url(payload.get("url", "")),
                "chat_title": payload.get("chatTitle") or payload.get("title") or "",
                "question": payload.get("question") or "",
                "page_url": payload.get("url") or "",
                "status": payload.get("status") or "",
                "complete": payload.get("complete"),
                "count": payload.get("count"),
                "expected_count": payload.get("expectedCount"),
                "index": item.get("index"),
                "title": item.get("title") or "",
                "href": href,
                "source": item.get("source") or "",
                "domain": domain_from_url(href),
                "extracted_at": payload.get("extractedAt") or "",
            })
    return run_no, len(rows)


def write_xlsx_from_csv():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except Exception:
        return False

    if not os.path.exists(OUT_CSV):
        return False

    with open(OUT_CSV, "r", encoding="utf-8-sig", newline="") as f:
        data = list(csv.reader(f))
    if not data:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "refs"
    for row in data:
        ws.append(row)

    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 8,
        "B": 20,
        "C": 42,
        "D": 28,
        "E": 36,
        "F": 70,
        "G": 12,
        "H": 10,
        "I": 8,
        "J": 14,
        "K": 8,
        "L": 80,
        "M": 90,
        "N": 28,
        "O": 28,
        "P": 26,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    end_row = max(ws.max_row, 2)
    end_col = ws.max_column
    end_letter = get_column_letter(end_col)
    table = Table(displayName="DeepSeekRefs", ref=f"A1:{end_letter}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)

    try:
        wb.save(OUT_XLSX)
        return True
    except PermissionError:
        return False


def main():
    payload = load_payload(sys.argv[1] if len(sys.argv) > 1 else sys.stdin.read())
    run_no, count = append_csv(payload)
    xlsx_ok = write_xlsx_from_csv()
    print(json.dumps({
        "ok": True,
        "run_no": run_no,
        "rows": count,
        "csv": OUT_CSV,
        "xlsx": OUT_XLSX,
        "xlsx_ok": xlsx_ok,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
