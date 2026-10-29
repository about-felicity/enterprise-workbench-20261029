from fastapi import APIRouter, Depends, Query, HTTPException, status, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
from openpyxl import Workbook
from datetime import datetime
from jose import JWTError, jwt

from database import get_db
from models import Device, Employee, Assignment, User
from auth import SECRET_KEY, ALGORITHM, oauth2_scheme

router = APIRouter(prefix="/api/export", tags=["数据导出"])


def get_user_export(token: str | None = Query(None), request: Request = None, db: Session = Depends(get_db)):
    """认证：优先从 query param token，其次从 Authorization header"""
    if token:
        try:
            payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
            uid = payload.get("sub")
            if uid:
                user = db.query(User).filter(User.id == int(uid)).first()
                if user and user.is_active:
                    return user
        except (JWTError, ValueError):
            pass
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="无效的token")

    # Try Bearer header
    auth_header = request.headers.get("Authorization") if request else None
    if auth_header and auth_header.startswith("Bearer "):
        try:
            token_str = auth_header.split(" ")[1]
            payload = jwt.decode(token_str, SECRET_KEY, algorithms=[ALGORITHM])
            uid = payload.get("sub")
            if uid:
                user = db.query(User).filter(User.id == int(uid)).first()
                if user and user.is_active:
                    return user
        except (JWTError, ValueError):
            pass

    raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="请先登录")


@router.get("/devices")
def export_devices(db: Session = Depends(get_db), current_user: User = Depends(get_user_export)):
    wb = Workbook()
    ws = wb.active
    ws.title = "设备清单"

    headers = ["设备编码", "设备类型", "品牌", "型号", "规格", "数量", "状态", "当前使用者", "采购日期", "保修截止", "备注"]
    ws.append(headers)

    devices = db.query(Device).order_by(Device.created_at.desc()).all()
    for d in devices:
        current_user_name = ""
        active = db.query(Assignment).filter(
            Assignment.device_id == d.id, Assignment.returned_at.is_(None)
        ).first()
        if active:
            emp = db.query(Employee).filter(Employee.id == active.employee_id).first()
            current_user_name = emp.name if emp else ""

        ws.append([
            d.device_code, d.device_type, d.brand, d.model, d.specs, d.quantity,
            d.status, current_user_name, d.purchase_date, d.warranty_expiry, d.notes,
        ])

    return _make_response(wb, "设备清单")


@router.get("/employees")
def export_employees(db: Session = Depends(get_db), current_user: User = Depends(get_user_export)):
    wb = Workbook()
    ws = wb.active
    ws.title = "员工清单"

    headers = ["姓名", "电话", "部门", "职位", "状态", "当前持有设备", "创建时间"]
    ws.append(headers)

    employees = db.query(Employee).order_by(Employee.created_at.desc()).all()
    for emp in employees:
        current_devices = ""
        active_assignments = db.query(Assignment).filter(
            Assignment.employee_id == emp.id, Assignment.returned_at.is_(None)
        ).all()
        for a in active_assignments:
            d = db.query(Device).filter(Device.id == a.device_id).first()
            if d:
                current_devices += f"{d.device_code}({d.device_type})；"

        ws.append([
            emp.name, emp.phone, emp.department, emp.position, emp.status,
            current_devices, emp.created_at.strftime("%Y-%m-%d") if emp.created_at else "",
        ])

    return _make_response(wb, "员工清单")


@router.get("/assignments")
def export_assignments(db: Session = Depends(get_db), current_user: User = Depends(get_user_export)):
    wb = Workbook()
    ws = wb.active
    ws.title = "分配记录"

    headers = ["设备编码", "设备类型", "员工姓名", "操作类型", "操作人", "分配时间", "归还时间", "备注"]
    ws.append(headers)

    records = db.query(Assignment).order_by(Assignment.assigned_at.desc()).all()
    for a in records:
        device = db.query(Device).filter(Device.id == a.device_id).first()
        employee = db.query(Employee).filter(Employee.id == a.employee_id).first()
        assigner = db.query(User).filter(User.id == a.assigned_by).first()

        ws.append([
            device.device_code if device else "", device.device_type if device else "",
            employee.name if employee else "", a.event_type,
            assigner.username if assigner else "",
            a.assigned_at.strftime("%Y-%m-%d %H:%M") if a.assigned_at else "",
            a.returned_at.strftime("%Y-%m-%d %H:%M") if a.returned_at else "",
            a.notes,
        ])

    return _make_response(wb, "分配记录")


def _make_response(wb: Workbook, name: str):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"{name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
