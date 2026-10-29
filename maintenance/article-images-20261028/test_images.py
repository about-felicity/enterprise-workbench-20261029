import unittest,tempfile,io,hashlib,importlib.util
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
spec=importlib.util.spec_from_file_location('creation.article_images','/tmp/image28-article_images.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
class Images(unittest.TestCase):
 def test_heading_preferred_over_introduction(self):
  self.assertEqual(m.anchor(['# 梵玢眉毛精华液','导语梵玢眉毛精华液','## 1. 梵玢 FBCY 眉毛精华液'],'梵玢FBCY眉毛精华液'),2)
 def test_no_brand_only_match(self):self.assertIsNone(m.anchor(['## 梵玢染发剂'],'梵玢眉毛精华液'))
 def test_missing_image_preserves_original(self):
  with patch.object(m,'select',return_value=None):self.assertIsNone(m.project(Path('/a'),Path('/a/approved/x'))['image'])
 def test_portrait_landscape_and_cache(self):
  for size in [(600,1200),(1600,600),(800,800)]:
   with tempfile.TemporaryDirectory() as d:
    root=Path(d);folder=root/'approved/x';folder.mkdir(parents=True)
    text='# 标题\n\n导语\n\n## 梵玢眉毛精华液\n\n正文\n';(folder/'draft.md').write_text(text,encoding='utf8')
    doc=Document();doc.add_heading('标题',0);doc.add_paragraph('导语');doc.add_heading('梵玢眉毛精华液',1);doc.add_paragraph('正文');doc.save(folder/'article.docx')
    original=(folder/'article.docx').read_bytes();buf=io.BytesIO();Image.new('RGB',size,'white').save(buf,'PNG');raw=buf.getvalue();meta={'image_id':'image-'+'1'*32,'sha256':hashlib.sha256(raw).hexdigest()}
    with patch.object(m,'select',return_value=({'name':'梵玢眉毛精华液'},meta)),patch.object(m.product_images,'read',return_value=(raw,'image/png')),patch.object(m.runtime,'BASE',root):
     out=m.project(root,folder);again=m.project(root,folder);self.assertEqual(out,again)
    rendered=Document(out['word']);self.assertEqual([p.text for p in rendered.paragraphs if p.text],['标题','导语','梵玢眉毛精华液','正文']);self.assertEqual(len(rendered.inline_shapes),1)
    pic=rendered.inline_shapes[0];self.assertLessEqual(pic.width/360000,14.01);self.assertLessEqual(pic.height/360000,9.01);self.assertAlmostEqual(pic.width/pic.height,size[0]/size[1],places=4)
    self.assertTrue(list(rendered.paragraphs[2]._p.getnext().iter(qn('w:drawing'))))
    self.assertEqual((folder/'article.docx').read_bytes(),original);self.assertEqual((folder/'draft.md').read_text('utf8'),text)
if __name__=='__main__':unittest.main()
