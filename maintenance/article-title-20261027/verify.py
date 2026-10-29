import hashlib,json,io,urllib.request,urllib.parse,zipfile
from pathlib import Path
from docx import Document
from creation import runtime
D=Path('/opt/monitor/geo-content/maintenance/20261027-article-title')
report=json.loads((D/'repair-1793067758.json').read_text());checked=[]
for row in report['changed']:
 folder=Path(row['path']);draft=(folder/'draft.md').read_text();old=(D/'articles-before'/folder.relative_to(runtime.BASE)/'draft.md').read_text()
 assert draft.endswith(old.strip()+'\n') or draft.split('\n\n',1)[1].strip()==old.strip().split('\n',1)[1].strip()
 title=draft.splitlines()[0][2:];assert len(title)<=80
 doc=Document(folder/'article.docx');assert doc.paragraphs[0].style.name=='Title'
 assert doc.paragraphs[0].text.replace('\n','')==title
 receipt=runtime.read(folder/'receipt.json')
 for name,digest in receipt['files'].items():assert hashlib.sha256((folder/name).read_bytes()).hexdigest()==digest
 checked.append(str(folder))
data=json.load(urllib.request.urlopen('http://127.0.0.1:8804/api/status',timeout=45))
root=Path(checked[0]).parents[3];folder=Path(checked[0]);article=folder.relative_to(root/'approved').as_posix()
downloads=[]
for kind in ('markdown','word','package'):
 url='http://127.0.0.1:8804/api/download?'+urllib.parse.urlencode({'job':root.name,'article':article,'kind':kind})
 response=urllib.request.urlopen(url,timeout=45);body=response.read()
 assert hashlib.sha256(body).hexdigest()==response.headers['X-Content-SHA256']
 if kind!='markdown':assert zipfile.ZipFile(io.BytesIO(body)).testzip() is None
 downloads.append({'kind':kind,'bytes':len(body),'status':response.status})
result={'repaired_verified':len(checked),'http_status':True,'downloads':downloads}
(D/'verification.json').write_text(json.dumps(result,indent=2));print(json.dumps(result))
