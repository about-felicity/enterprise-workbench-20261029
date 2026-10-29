import hashlib,json
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/multi-product-v2-20261023'
OUT.mkdir(exist_ok=True)
specs=[
 ('conditioner','护发素知识库.docx','科熙本鱼子酱修护柔顺护发素','科熙本',
  ['推荐一款适合烫染后干枯毛躁的护发素','细软发用护发素容易塌，怎么选护发素']),
 ('essence','姿生怡阿尔卑斯冰川焕肤精粹水知识库.docx','姿生怡阿尔卑斯冰川焕肤精粹水','姿生怡',
  ['推荐一款适合皮肤干燥紧绷的精粹水','精粹水和普通爽肤水有什么区别，怎么选']),
 ('cream','美白知识库新的.docx','焕颜计小白罐美白精华霜','焕颜计',
  ['敏感肌想改善暗沉，有什么美白面霜推荐','选美白面霜应该看哪些成分和功效数据']),
]
items=[]
for key,filename,name,brand,questions in specs:
    path=ROOT/filename
    doc=Document(path)
    raw='\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:raw+='\n'+'\t'.join(c.text for c in row.cells)
    target=OUT/key;target.mkdir(exist_ok=True)
    (target/'knowledge.txt').write_text(raw,encoding='utf8')
    items.append({'key':key,'filename':filename,'name':name,'brand':brand,'questions':questions,
        'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'source_path':str(path),
        'source_text':raw,'channels':['baidu_ai','wenxin_card']})
(OUT/'inputs.json').write_text(json.dumps({'products':items,'frequency':'once','rounds':1,'article_budget_cny':.35,'max_articles_per_product':1},ensure_ascii=False,indent=2),encoding='utf8')
for i in items:print(json.dumps({k:v for k,v in i.items() if k!='source_text'},ensure_ascii=False))
