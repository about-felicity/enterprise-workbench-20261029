"""Zero-call input propagation check using the preserved real source and Word.

The writer answer is a previously saved draft, NOT a new model generation.
The skeleton approval is a test fixture in this isolated script only.
"""
import json
import hashlib
import sys
from pathlib import Path
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / 'geo-source-card-collectors'))
from geo_source_cards.rewrite_v2 import full_writer, material_manager, structure_parser
from geo_source_cards.rewrite_v2.common import seal, digest

original = ROOT / 'outputs/rank01-readable-draft-20261023-v2'
sources = json.loads((original / 'product_data_sources.json').read_text('utf8'))
word = Path(sources['source_file'])
assert hashlib.sha256(word.read_bytes()).hexdigest() == sources['source_sha256']
document = Document(word)
raw = '\n'.join(p.text for p in document.paragraphs if p.text.strip())
for table in document.tables:
    for row in table.rows:
        raw += '\n' + '\t'.join(c.text for c in row.cells)
product = json.loads((original / 'inputs/product.json').read_text('utf8'))
product = {**product, 'raw_source_text': raw, 'raw_source_sha256': sources['source_sha256']}
inputs = json.loads((original / 'inputs/research_input.json').read_text('utf8'))
url = inputs['reference_design']['primary_url']
article = next(s['article'] for r in inputs['observations'] for s in r['sources'] if s['url'] == url)
reference = {'url': url, 'title': article['title'], 'text': article['text']}
skeleton = {'article_type': 'ranking', 'tone': '验证用编辑视角',
            'article_structure': ['读者选购问题', '比较标准', '逐款介绍', '横向总结与边界'],
            'writing_rhythm': '问题到比较再到总结', 'special_style_elements': [], 'forbidden_pattern': []}
envelope = seal({'protocol': structure_parser.PROTOCOL, 'cache_key': structure_parser.identity(reference['text']),
                 'source_text': reference['text'], 'skeleton': skeleton,
                 'approval': {'reviewer': 'OFFLINE_TEST_FIXTURE_NOT_HUMAN_APPROVAL', 'skeleton_sha256': digest(skeleton)}})
material = seal({'protocol': material_manager.PROTOCOL, 'product_name': product['name'],
                 'raw_product': product, 'layers': {k: [] for k in material_manager.LAYERS}, 'usage_rules': material_manager.RULES})
context = {'workflow': 'rewrite-v2', 'reference': reference, 'skeleton': envelope,
           'product': product, 'category': {'label': '排行类', 'number': 1}}

class Capture:
    def complete(self, prompt, directory, **kwargs):
        self.prompt = prompt
        return {'draft_markdown': (original / 'draft.md').read_text('utf8')}

client = Capture()
answer = full_writer.generate(client, context, material, HERE / 'unused', {'date': '2026-10-23', 'category': '排行类', 'number': 1})
data = json.loads(client.prompt[client.prompt.index('{"skeleton":'):])
assert data['material']['raw_product']['raw_source_text'] == raw
assert data['reference']['text'] == reference['text']
checks = {key: key in client.prompt for key in ['43.96%', '33.96%', '27.53%', '14200', 'GT00332022056002', 'G20211805', '16.0%', '140mm', '5.9', '158']}
assert all(checks.values()), checks
out = ROOT / 'outputs/rewrite-v2-validation-20261023'
out.mkdir(exist_ok=True)
(out / 'existing-draft-roundtrip.md').write_text(answer['draft_markdown'], encoding='utf8')
report = {'real_model_calls': 0, 'new_generated_article': False,
          'fixture_skeleton_not_production_approval': True, 'source_word_sha256': sources['source_sha256'],
          'source_word_text_characters': len(raw), 'reference_characters': len(reference['text']),
          'all_word_text_preserved_in_writer_input': True, 'all_reference_text_preserved_in_writer_input': True,
          'numeric_and_report_checks': checks, 'writer_prompt_utf8_bytes': len(client.prompt.encode()),
          'word_roundtrip_input': str(out / 'existing-draft-roundtrip.md'),
          'article_quality_verified': False, 'runtime_and_cost_target_verified': False}
(HERE / 'real-material-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps(report, ensure_ascii=False, indent=2))
