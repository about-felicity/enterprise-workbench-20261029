"""Freeze the user-provided Word and a previously captured reference for a new sample."""
import hashlib
import json
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/rewrite-v2-real-sample-20261023'
OUT.mkdir(exist_ok=True)
saved = ROOT / 'outputs/rank01-readable-draft-20261023-v2/inputs'
inputs = json.loads((saved / 'research_input.json').read_text('utf8'))
source = next(s for r in inputs['observations'] for s in r['sources'] if '1879562453794511892' in s['url'])
word = Path('C:/Users/78567/Downloads/2.道和小绿瓶防脱洗发水知识库(1).docx')
doc = Document(word)
raw = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
for table in doc.tables:
    for row in table.rows:
        raw += '\n' + '\t'.join(cell.text for cell in row.cells)
product = json.loads((saved / 'product.json').read_text('utf8'))
product.update(raw_source_text=raw, raw_source_sha256=hashlib.sha256(word.read_bytes()).hexdigest())
data = {'product': product, 'observations': [{'question': '防脱洗发水哪个牌子的效果好', 'sources': [source]}],
        'provenance': {'word_path': str(word), 'word_sha256': product['raw_source_sha256'],
                       'saved_research_input': str(saved / 'research_input.json'),
                       'reference_url': source['url'], 'capture_day': '2026-10-21',
                       'new_fetch': False, 'new_sample': True}}
(OUT / 'sample_input.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf8')
(OUT / 'reference.md').write_text('# ' + source['article']['title'] + '\n\n' + source['url'] + '\n\n' + source['article']['text'], encoding='utf8')
(OUT / 'knowledge-source.txt').write_text(raw, encoding='utf8')
print(json.dumps({'path': str(OUT), 'url': source['url'], 'reference_chars': len(source['article']['text']), 'knowledge_chars': len(raw)}, ensure_ascii=False))
