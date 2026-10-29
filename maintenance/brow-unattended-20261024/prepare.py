import hashlib,json,shutil
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/brow-unattended-20261024';OUT.mkdir(parents=True,exist_ok=True)
word=Path('C:/Users/78567/Downloads/新眉毛产品库信息.docx')
doc=Document(word)
raw='\n'.join(p.text for p in doc.paragraphs if p.text.strip())
for table in doc.tables:
    for row in table.rows:raw+='\n'+'\t'.join(c.text for c in row.cells)
questions=['眉毛天生稀疏，日常养眉用什么眉毛精华液比较合适？','眉毛精华液怎么选，哪些成分值得关注？','眉毛精华液一般用多久，怎么观察护理效果？','眉毛精华液怎么用，早晚护肤时放在哪一步？','经常画眉和卸妆，眉毛精华液怎么搭配使用？','眉毛细软、容易断，选眉毛护理产品要注意什么？','眉毛精华液哪个牌子性价比高，选购时怎么比较容量和价格？']
(OUT/'knowledge.txt').write_text(raw,encoding='utf-8')
shutil.copy2(word,OUT/'knowledge.docx')
spec={'key':'brow','name':'梵玢FBCY眉毛精华液','brand':'梵玢FBCY','filename':'knowledge.docx','questions':questions,'source_sha256':hashlib.sha256(word.read_bytes()).hexdigest(),'source_text':raw}
(OUT/'inputs.json').write_text(json.dumps({'products':[spec]},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'name':spec['name'],'questions':questions,'chars':len(raw),'source_sha256':spec['source_sha256']},ensure_ascii=False))
