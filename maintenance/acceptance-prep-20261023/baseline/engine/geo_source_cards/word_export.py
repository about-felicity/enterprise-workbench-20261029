"""Portable article-only Word export; quality/provenance remains in sidecar JSON."""

from __future__ import annotations

import argparse
import re
from pathlib import Path


def export_docx(markdown, output: Path, *, demo=False, typography=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(2.3)
    section.left_margin = section.right_margin = Cm(2.5)
    typography = typography or {
        "chinese_font": "宋体",
        "latin_font": "Cambria",
        "body_size": 11,
        "title_size": 20,
        "heading_size": 14,
        "line_spacing": 1.5,
    }
    # Explicit fonts prevent Word/WPS from selecting a Japanese theme fallback.
    for element in document.styles.element.iter(qn("w:rFonts")):
        for key in list(element.attrib):
            if "theme" in key.lower():
                del element.attrib[key]
        element.set(qn("w:ascii"), typography["latin_font"])
        element.set(qn("w:hAnsi"), typography["latin_font"])
        element.set(qn("w:eastAsia"), typography["chinese_font"])
    for element in list(document.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    for name, size in [
        ("Normal", typography["body_size"]),
        ("Title", typography["title_size"]),
        ("Heading 1", typography["heading_size"]),
        ("Heading 2", 12),
        ("Heading 3", 11),
        ("Heading 4", 11),
        ("Heading 5", 11),
    ]:
        style = document.styles[name]
        style.font.name = typography["latin_font"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), typography["chinese_font"])
        style.paragraph_format.space_before = Pt(
            10 if name.startswith("Heading") else 0
        )
        style.paragraph_format.space_after = Pt(6 if name == "Normal" else 10)
        style.paragraph_format.line_spacing = typography["line_spacing"]
        style.font.bold = name != "Normal"
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True
        style.paragraph_format.first_line_indent = Pt(2 * typography["body_size"] if name == "Normal" else 0)
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles["Title"].paragraph_format.line_spacing = 1.25
    document.styles["Title"].paragraph_format.space_after = Pt(18)

    def add_text(text, style=None):
        paragraph = document.add_paragraph(style=style)
        # Preserve inline emphasis instead of flattening it into ordinary text.
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
            if part.startswith("**") and part.endswith("**"):
                run.bold = True
        return paragraph

    title_written = False
    for block in re.split(r"\n\s*\n", markdown.strip()):
        for line in block.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith("# "):
                title = line[2:].strip()
                width = sum(1 if ord(c) > 127 else 0.5 for c in title)
                if width > 22:
                    breaks = [m.end() for m in re.finditer(r"[：:，,？?！!]", title)
                              if len(title) * .25 < m.end() < len(title) * .75]
                    if breaks:
                        split = min(breaks, key=lambda n: abs(n - len(title) / 2))
                        title = title[:split] + "\n" + title[split:]
                add_text(
                    title, style="Title" if not title_written else "Heading 1"
                )
                title_written = True
                if demo:
                    document.add_paragraph("示例稿 非正式产品宣传材料")
            elif line.startswith("## "):
                add_text(line[3:].strip(), style="Heading 1")
            elif re.match(r"^#{3,6}\s",line):
                marker,text=line.split(maxsplit=1)
                add_text(text.strip(), style="Heading "+str(len(marker)-1))
            elif re.match(r"^(?:[-*+] |\d+[.)、]\s*)", line):
                # Keep source numbering exactly; hanging indent aligns wrapped lines.
                text = re.sub(r"^[-*+] ", "• ", line)
                paragraph = add_text(text)
                paragraph.paragraph_format.left_indent = Pt(24)
                paragraph.paragraph_format.first_line_indent = Pt(-24)
            else:
                add_text(line)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    staged = output.with_name(output.stem + ".tmp.docx")
    document.save(staged)
    staged.replace(output)
    return output


def main(argv=None):
    p = argparse.ArgumentParser(
        description="Export a complete Markdown article to Word"
    )
    p.add_argument("--input", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--demo", action="store_true")
    a = p.parse_args(argv)
    export_docx(a.input.read_text(encoding="utf-8-sig"), a.output, demo=a.demo)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
