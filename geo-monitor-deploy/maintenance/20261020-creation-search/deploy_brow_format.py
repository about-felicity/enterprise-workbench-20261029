"""Re-export unchanged reviewed Markdown; preserve original provider artifacts."""
import hashlib,json,re,shutil,subprocess,sys,time
from pathlib import Path
base=Path('/opt/monitor/geo-content')
def read(p):return json.loads(p.read_text()) if p.exists() else {}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
t=read(Path('/opt/monitor/geo-capture/maintenance/20261020-brow-daily/task.json'))
roots=[p.parent for p in (base/'runtime/runs').glob('*/job.json') if read(p).get('studio_plan_id')==t['plan_id']];assert len(roots)==1
r=roots[0];job=read(r/'job.json');day=job['date']
assert read(r/'days'/day/'day.json').get('status')=='complete'
assert not read(r/'control.json').get('paused')
for p in (base/'runtime/runs').glob('*/job.json'):
    if read(p.parent/'control.json').get('paused'):continue
    assert read(p.parent/'days'/read(p)['date']/'day.json').get('status')=='complete',str(p)
backup=Path('/opt/monitor/geo-capture/maintenance/20261020-brow-daily/before-word-headings-v1')
assert not backup.exists(),'Already deployed; inspect existing formatting report'
targets={'/tmp/brow-word-export.py':base/'engine/geo_source_cards/word_export.py','/tmp/brow-article-export.py':base/'creation/article_export.py'}
for source,target in targets.items():compile(Path(source).read_bytes(),str(target),'exec')
subprocess.run(['systemctl','stop','geo-content-worker.service'],check=True)
try:
    backup.mkdir(parents=True,mode=0o700)
    for source,target in targets.items():
        shutil.copy2(target,backup/target.name)
        temp=target.with_suffix('.heading-new');temp.write_bytes(Path(source).read_bytes());temp.replace(target)
    sys.path[:0]=[str(base),str(base/'engine')]
    from creation import runtime
    from docx import Document
    from geo_source_cards.word_export import export_docx
    from geo_source_cards.pipeline_control import event
    changed=[]
    for folder in sorted((r/'approved'/day).glob('*/*')):
        receipt=runtime.read(folder/'receipt.json');word=folder/'article.docx'
        if not receipt:continue
        assert sha(word)==receipt['word_sha256']
        before_text=[p.text for p in Document(word).paragraphs]
        if not any(re.match(r'^#{4,6}\s',p) for p in before_text):continue
        old=r/'repair_evidence/word-headings-v1'/folder.parent.name/folder.name
        old.mkdir(parents=True,exist_ok=False)
        for name in ('article.docx','receipt.json','draft.md','semantic_review.json'):shutil.copy2(folder/name,old/name)
        draft_hash=sha(folder/'draft.md');new_word=folder/'formatted.docx'
        export_docx((folder/'draft.md').read_text(),new_word)
        expected=[re.sub(r'^#{4,6}\s+','',p) for p in before_text]
        assert [p.text for p in Document(new_word).paragraphs]==expected,'Formatting changed article text'
        assert sha(folder/'draft.md')==draft_hash
        new_sha=sha(new_word);new_word.replace(word)
        note={'at':time.time(),'reason':'Render Markdown heading levels 4-6 as black Word headings instead of literal # markers','draft_sha256_unchanged':draft_hash,'original_word_sha256':receipt['word_sha256'],'formatted_word_sha256':new_sha,'original_receipt_sha256':sha(old/'receipt.json'),'provider_artifacts_unchanged':True,'new_api_calls':0,'content_review_unchanged':True}
        runtime.atomic_json(folder/'formatting_revision.json',note)
        files={**receipt['files'],'article.docx':new_sha,'formatting_revision.json':sha(folder/'formatting_revision.json')}
        runtime.atomic_json(folder/'receipt.json',{**receipt,'word_sha256':new_sha,'files':files,'formatting_revision':'word-headings-v1'})
        event(r,'formatting','从同一已审稿Markdown修正Word深层子标题，正文与内容审稿不变，未调用API',category=receipt['category'],number=receipt['number'])
        changed.append({'article':str(folder.relative_to(r/'approved')),'before':note['original_word_sha256'],'after':new_sha})
    runtime.atomic_json(Path('/opt/monitor/geo-capture/maintenance/20261020-brow-daily/word-formatting.json'),{'at':time.time(),'changed':changed,'new_api_calls':0})
finally:subprocess.run(['systemctl','start','geo-content-worker.service'],check=True)
subprocess.run(['systemctl','restart','geo-content-api.service'],check=True)
print(json.dumps({'formatted_articles':changed,'new_api_calls':0},ensure_ascii=False))
